"""Build verified historical win-rate/PnL context for the public basketball page.

The workbook is a human-maintained paper ledger.  This reader never changes the
workbook or the immutable JSON ledgers; it only emits a small, explicit display
index.  Exact same-pair history is kept separate from same-league team history
so the dashboard never borrows an unrelated bucket silently.
"""
from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


def clean_team(name: str) -> str:
    text = str(name or '').strip()
    text = re.sub(r'\[[^\]]*\]', '', text)
    text = re.sub(r'\s+', '', text)
    return text


def league_family(value: str) -> str:
    """Collapse phase labels so NBA preseason/regular/postseason share NBA."""
    text = re.sub(r'\s+', '', str(value or '').strip())
    # Source labels vary between full and abbreviated phase names.
    text = re.sub(r'(季前赛?|常规赛?|季后赛?|附加赛?|资格赛?|淘汰赛?|半决赛|总决赛|决赛)$', '', text)
    return text or str(value or '').strip()


def split_match(value: str):
    parts = re.split(r'\s+vs\s+', str(value or ''), maxsplit=1)
    if len(parts) != 2:
        return '', ''
    return clean_team(parts[0]), clean_team(parts[1])


def stat(rows):
    settled = [r for r in rows if r['outcome'] in {'赢', '输', '走'}]
    wins = sum(r['outcome'] == '赢' for r in settled)
    losses = sum(r['outcome'] == '输' for r in settled)
    pushes = sum(r['outcome'] == '走' for r in settled)
    pnl = round(sum(r['pnl'] for r in settled), 4)
    n = len(settled)
    return {
        'n': n, 'wins': wins, 'losses': losses, 'pushes': pushes,
        'win_rate': round(wins / n, 4) if n else None,
        'pnl_units': pnl if n else None,
        'roi': round(pnl / n, 4) if n else None,
        'status': 'SETTLED_SAMPLE' if n else 'NO_EXACT_SAMPLE',
    }


def read_ledger(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb['赛事明细']
    rows = []
    for values in ws.iter_rows(min_row=4, values_only=True):
        values = list(values) + [None] * 13
        if values[12] != '计入' or values[8] not in {'赢', '输', '走'}:
            continue
        t1, t2 = split_match(values[3])
        if not t1 or not t2:
            continue
        try:
            pnl = float(values[10] or 0)
        except (TypeError, ValueError):
            pnl = 0.0
        rows.append({
            'date': str(values[0] or ''),
            'league': str(values[2] or ''),
            'league_family': league_family(values[2] or ''),
            'market': str(values[4] or ''),
            'team1': t1, 'team2': t2,
            'outcome': str(values[8]), 'pnl': pnl,
        })
    return rows


def build(workbook_path: Path, current_path: Path):
    history = read_ledger(workbook_path)
    current = json.loads(current_path.read_text(encoding='utf-8'))
    records = {}
    for item in current.get('events', []):
        if not (item.get('execution') or {}).get('triggered'):
            continue
        event = item.get('event') or {}
        t1, t2 = clean_team(event.get('team1_name')), clean_team(event.get('team2_name'))
        league = str(event.get('league_display') or event.get('league') or event.get('competition_key') or '')
        family = league_family(league)
        market = '让分'
        exact = [r for r in history if r['market'] == market and r['league_family'] == family and {r['team1'], r['team2']} == {t1, t2}]
        team1 = [r for r in history if r['market'] == market and r['league_family'] == family and t1 in {r['team1'], r['team2']}]
        team2 = [r for r in history if r['market'] == market and r['league_family'] == family and t2 in {r['team1'], r['team2']}]
        league_rows = [r for r in history if r['market'] == market and r['league_family'] == family]
        records[str(event.get('event_id') or item.get('id'))] = {
            'event_id': str(event.get('event_id') or item.get('id')),
            'league': league, 'league_family': family, 'team1': event.get('team1_name'), 'team2': event.get('team2_name'),
            'same_matchup': stat(exact), 'league_overall': stat(league_rows),
            'team1_same_league': stat(team1), 'team2_same_league': stat(team2),
            'source': str(workbook_path),
        }
    return {
        'schema_version': 'paper-ledger-history-1.0',
        'generated_at': datetime.now(timezone.utc).isoformat(),
        'source_workbook': str(workbook_path),
        'settled_ledger_rows': len(history),
        'records': records,
        'display_note': '联赛统计按联赛主名合并季前、常规赛、季后赛；仅使用累计 Excel 中已结算且统计纳入的同市场记录。',
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--workbook', required=True)
    ap.add_argument('--current-json', required=True)
    ap.add_argument('--output', required=True)
    args = ap.parse_args()
    result = build(Path(args.workbook), Path(args.current_json))
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    Path(args.output).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'output': args.output, 'records': len(result['records']), 'settled_rows': result['settled_ledger_rows']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
