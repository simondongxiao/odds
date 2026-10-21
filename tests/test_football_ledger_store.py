from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, r"D:\codex\tools")
import football_ledger_store as store
import update_football_history_ledgers as legacy


HEADERS = [
    "日期筛选", "比赛时间", "赛事分布", "赛事类型", "联赛/杯赛", "具体比赛",
    "即时主水", "即时盘口", "即时客水", "即时赔率文本", "V4等级/信号", "投注方",
    "模型有效胜率", "投注说明", "赛果", "结算状态", "红黑", "投注水位", "盈亏(u)",
    "来源", "结算水位", "比分来源", "结算盈亏(u)", "Match_ID",
]


def make_book(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "赛事明细"
    for column, header in enumerate(HEADERS, 1):
        ws.cell(3, column).value = header
    values = ["2026-10-01", "2026-10-01 20:00:00", "欧洲", "联赛", "意丙", "甲队 vs 乙队", 0.8, -0.5, 1.0, "0.8/-0.5/1.0", "V4-B", "甲队", 0.6, "frozen", "1-0", "已结算", "红", 0.8, 0.8, "old", None, "source", None, "m1"]
    for column, value in enumerate(values, 1):
        ws.cell(4, column).value = value
    second = ["2026-10-01", "2026-10-01 21:00:00", "欧洲", "联赛", "意丙", "己队 vs 庚队", 0.9, "0", 0.9, "0.9/0/0.9", "V4-N", "", 0.5, "keep-me", None, "未开赛", None, 0.9, None, "old", None, None, None, "m2"]
    for column, value in enumerate(second, 1):
        ws.cell(5, column).value = value
    other = wb.create_sheet("其他sheet")
    other["A1"] = "格式和其他sheet必须保留"
    table = Table(displayName="MatchLedger", ref=f"A3:X5")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    wb.save(path)


def test_started_or_settled_row_preserves_prematch_fields(tmp_path: Path) -> None:
    book = tmp_path / "ledger.xlsx"
    make_book(book)
    raw = pd.DataFrame([{
        "match_id": "m1", "list_date": "2026-10-01", "bj_time": "10-01 20:00",
        "home_cn": "甲队", "away_cn": "乙队", "league_cn": "意丙", "state": "-1",
        "home_score": "2", "away_score": "1", "ah_full_current_home_or_over": "0.99",
        "ah_full_current_line_or_draw": "-0.75", "ah_full_current_away_or_under": "0.88",
    }])
    raw["_unique_key"] = raw.apply(lambda row: store._incoming_key(row.to_dict()), axis=1)
    model = {"m1": {"match_id": "m1", "grade": "A", "analysis_status": "EVALUATED", "selected_team": "甲队", "quality_gate_passed": True, "final_decision": "BET_GIVING"}}
    store.update_workbook(book, "V4", raw, model, {}, "new-source", tmp_path / "backup", cache_path=tmp_path / "cache.json")
    wb = load_workbook(book, data_only=False)
    ws = wb["赛事明细"]
    assert ws["G4"].value == 0.8
    assert ws["H4"].value == -0.5
    assert ws["N4"].value == "frozen"
    assert ws["O4"].value == "2-1"
    assert ws["P4"].value == "已结算"
    assert ws["F5"].value == "己队 vs 庚队"
    assert ws["N5"].value == "keep-me"
    assert "其他sheet" in wb.sheetnames
    assert wb["其他sheet"]["A1"].value == "格式和其他sheet必须保留"
    assert ws.tables["MatchLedger"].ref.endswith("5")
    wb.close()


def test_history_fields_and_three_way_reconciliation(tmp_path: Path) -> None:
    rows = [{"日期筛选": "2026-10-11", "联赛/杯赛": "意丙", "具体比赛": "H1 vs A1", "Match_ID": "1", "结算状态": "已结算", "红黑": "红", "盈亏(u)": 0.8, "Final_Decision": "BET_GIVING", "Row_Lock_State": "STARTED_OR_SETTLED_LOCKED"}]
    cache = tmp_path / "cache.json"; v4 = tmp_path / "v4.json"; warning = tmp_path / "warning.json"; workbook = tmp_path / "book.xlsx"
    cache.write_text(json.dumps({"rows": rows}, ensure_ascii=False), encoding="utf-8")
    v4.write_text(json.dumps({"matches": [{"match_id": "1", "quality_gate_passed": True, "final_decision": "BET_GIVING"}]}, ensure_ascii=False), encoding="utf-8")
    warning.write_text(json.dumps({"status": "PASS"}, ensure_ascii=False), encoding="utf-8")
    workbook.write_bytes(b"workbook-hash-fixture")
    result = store.export_dashboard(cache, v4, "2026-10-11", tmp_path / "dashboard", warning, workbook)
    assert result["reconciliation"]["status"] == "PASS"
    full = json.loads((tmp_path / "dashboard" / "data_full.json").read_text(encoding="utf-8"))
    assert full["rows"][0]["history_key"] == "意丙|id:1"
    assert full["rows"][0]["countable_history"] is True
    assert full["rows"][0]["frozen_actionable"] is True
    assert json.loads((tmp_path / "dashboard" / "three_way_reconciliation.json").read_text(encoding="utf-8"))["status"] == "PASS"


def test_three_way_reconciliation_blocks_tampered_full_json(tmp_path: Path) -> None:
    cache = tmp_path / "cache.json"; v4 = tmp_path / "v4.json"; warning = tmp_path / "warning.json"; workbook = tmp_path / "book.xlsx"
    cache.write_text(json.dumps({"rows": [{"日期筛选": "2026-10-11", "联赛/杯赛": "意丙", "具体比赛": "H1 vs A1", "Match_ID": "1", "结算状态": "未开赛"}]}, ensure_ascii=False), encoding="utf-8")
    v4.write_text(json.dumps({"matches": []}, ensure_ascii=False), encoding="utf-8")
    warning.write_text(json.dumps({"status": "PASS"}, ensure_ascii=False), encoding="utf-8")
    workbook.write_bytes(b"workbook-hash-fixture")
    output = tmp_path / "dashboard"
    store.export_dashboard(cache, v4, "2026-10-11", output, warning, workbook)
    full_path = output / "data_full.json"
    payload = json.loads(full_path.read_text(encoding="utf-8")); payload["rows"] = []
    full_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    audit = store._reconcile_views(json.loads(cache.read_text(encoding="utf-8"))["rows"], payload, json.loads((output / "data_filtered.json").read_text(encoding="utf-8")), output / "three_way_reconciliation.json", workbook)
    assert audit["status"] == "FAIL"
    assert "EXCEL_FULL_ROW_COUNT_MISMATCH" in audit["mismatches"]


def test_count_shrink_guard_blocks_43_to_16(tmp_path: Path) -> None:
    full = pd.DataFrame([{"match_id": str(i), "list_date": "2026-10-11", "league_cn": "意丙", "home_cn": f"H{i}", "away_cn": f"A{i}"} for i in range(43)])
    full["_unique_key"] = full.apply(lambda row: store._incoming_key(row.to_dict()), axis=1)
    contract, warning = tmp_path / "contract.json", tmp_path / "warning.json"
    assert store.audit_roster(full, "2026-10-11", contract, warning)["contract_total"] == 43
    with pytest.raises(store.RowCountShrinkError):
        store.audit_roster(full.iloc[:16].copy(), "2026-10-11", contract, warning)
    payload = json.loads(warning.read_text(encoding="utf-8"))
    assert payload["status"] == "ROW_COUNT_SHRINK_WARNING"
    assert payload["previous_total"] == 43
    assert payload["incoming_total"] == 16
    assert json.loads(contract.read_text(encoding="utf-8"))["total"] == 43
    exceptions = tmp_path / "exceptions.json"
    exceptions.write_text(json.dumps({"exceptions": [{"match_id": str(index), "status": "POSTPONED", "evidence": f"official-{index}"} for index in range(16, 43)]}), encoding="utf-8")
    allowed = store.audit_roster(full.iloc[:16].copy(), "2026-10-11", contract, warning, exceptions)
    assert allowed["status"] == "PASS"
    assert allowed["official_exception_count"] == 27


def test_dual_json_keeps_43_rows_and_uses_16_settled_samples(tmp_path: Path) -> None:
    rows = []
    models = []
    for index in range(43):
        settled = index < 16
        rows.append({"日期筛选": "2026-10-11", "联赛/杯赛": "意丙", "具体比赛": f"H{index} vs A{index}", "Match_ID": str(index), "结算状态": "已结算" if settled else "赛果未匹配待人工核验", "红黑": "红" if settled else "", "盈亏(u)": 0.8 if settled else None, "Frozen_Actionable": settled})
        # A quality-valid NO_BET row is still not part of the strategy view.
        models.append({"match_id": str(index), "quality_gate_passed": settled or index == 16, "final_decision": "BET_GIVING" if settled else "NO_BET", "grade": "B" if settled else "N", "analysis_status": "EVALUATED"})
    cache = tmp_path / "cache.json"; v4 = tmp_path / "v4.json"; warning = tmp_path / "warning.json"; workbook = tmp_path / "book.xlsx"
    cache.write_text(json.dumps({"rows": rows}, ensure_ascii=False), encoding="utf-8")
    v4.write_text(json.dumps({"matches": models}, ensure_ascii=False), encoding="utf-8")
    warning.write_text(json.dumps({"status": "PASS"}, ensure_ascii=False), encoding="utf-8")
    workbook.write_bytes(b"workbook-hash-fixture")
    result = store.export_dashboard(cache, v4, "2026-10-11", tmp_path / "dashboard", warning, workbook)
    assert result["full_rows"] == 43
    assert result["filtered_rows"] == 16
    full = json.loads((tmp_path / "dashboard" / "data_full.json").read_text(encoding="utf-8"))
    assert full["competition_stats"]["意丙"]["total_ledger_rows"] == 43
    assert full["competition_stats"]["意丙"]["settled_sample_rows"] == 16


def test_daily_type_mapping_does_not_copy_league_into_type() -> None:
    raw = {
        "list_date": "2026-10-09", "bj_time": "10-09 20:00", "league_cn": "英锦赛",
        "league_tw": "英锦赛", "home_cn": "甲队", "away_cn": "乙队", "match_id": "m-type",
        "state": "0", "home_score": "", "away_score": "",
    }
    values = legacy.v3_values(raw, {"Competition_Domain": "CLUB_CUP", "action": "不投"}, "source")
    assert values[3] == "杯赛"
    assert values[3] != values[4]


def test_v4_result_styles_keep_red_black_distinct(tmp_path: Path) -> None:
    book = tmp_path / "styles.xlsx"
    wb = Workbook(); ws = wb.active; ws.title = "赛事明细"
    for column, header in enumerate(HEADERS, 1):
        ws.cell(3, column).value = header
    for row, label in enumerate(["红", "黑", "半红", "半黑", "走"], 4):
        ws.cell(row, 1).value = "2026-10-08"
        ws.cell(row, 5).value = "意丙"
        ws.cell(row, 17).value = label
    table = Table(displayName="StyleLedger", ref="A3:X8")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table); wb.save(book); wb.close()
    wb = load_workbook(book); ws = wb["赛事明细"]; headers = store._headers(ws)
    assert store._normalize_v4_result_styles(ws, headers)
    fills = {ws.cell(row, 17).value: ws.cell(row, 17).fill.fgColor.rgb for row in range(4, 9)}
    assert fills["红"] != fills["黑"]
    assert fills["半红"] != fills["半黑"]
    wb.close()
