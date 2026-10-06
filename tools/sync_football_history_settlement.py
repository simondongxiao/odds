"""Write frozen V3/V4 settlement results back to the Excel history ledgers.

The daily HTML history is intentionally downstream of these workbooks. This
script updates only settlement fields for rows already present in the ledgers;
it never changes the frozen decision, grade, line, water, or probability.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import shutil
from pathlib import Path

from openpyxl import load_workbook


RESULT_LABEL = {"W": "红", "HW": "半红", "P": "走", "HL": "半黑", "L": "黑"}


def clean(value: object) -> str:
    return "" if value is None else " ".join(str(value).split()).strip()


def key_text(value: object) -> str:
    return clean(value).replace("（", "(").replace("）", ")")


def number(value: object) -> object:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def read_settlements(path: Path, list_date: str, version: str) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    return [
        row for row in rows
        if row.get("version") == version
        and row.get("list_date") == list_date
        and row.get("settlement_status") == "SETTLED"
        and row.get("result") in RESULT_LABEL
    ]


def headers(ws) -> dict[str, int]:
    return {clean(ws.cell(3, col).value): col for col in range(1, ws.max_column + 1) if ws.cell(3, col).value}


def update_ledger(path: Path, rows: list[dict[str, str]], version: str, list_date: str, backup_dir: Path) -> dict[str, object]:
    if not rows:
        return {"file": str(path), "rows_available": 0, "rows_updated": 0, "unmatched": []}
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / path.name
    shutil.copy2(path, backup)
    wb = load_workbook(path, data_only=False)
    ws = wb["赛事明细"]
    h = headers(ws)
    required = ["日期筛选", "具体比赛", "赛果", "结算状态", "红黑", "盈亏(u)"]
    missing = [name for name in required if name not in h]
    if missing:
        wb.close()
        raise RuntimeError(f"{path} missing columns: {missing}")
    date_col, match_col = h["日期筛选"], h["具体比赛"]
    id_col = h.get("Match_ID") or h.get("match_id")
    by_id = {clean(row.get("match_id")): row for row in rows if clean(row.get("match_id"))}
    by_match = {key_text(row.get("match")): row for row in rows if key_text(row.get("match"))}
    updated = 0
    matched_keys: set[str] = set()
    for rownum in range(4, ws.max_row + 1):
        if clean(ws.cell(rownum, date_col).value)[:10] != list_date:
            continue
        settlement = None
        if id_col:
            settlement = by_id.get(clean(ws.cell(rownum, id_col).value))
        if settlement is None:
            settlement = by_match.get(key_text(ws.cell(rownum, match_col).value))
        if settlement is None:
            continue
        ws.cell(rownum, h["赛果"]).value = settlement.get("score", "")
        ws.cell(rownum, h["结算状态"]).value = "已结算"
        ws.cell(rownum, h["红黑"]).value = RESULT_LABEL[settlement["result"]]
        ws.cell(rownum, h["盈亏(u)"]).value = number(settlement.get("pnl_1u"))
        updated += 1
        matched_keys.add(clean(settlement.get("match_id")) or key_text(settlement.get("match")))
    wb.save(path)
    wb.close()
    all_keys = {clean(row.get("match_id")) or key_text(row.get("match")) for row in rows}
    return {
        "file": str(path),
        "rows_available": len(rows),
        "rows_updated": updated,
        "unmatched": sorted(all_keys - matched_keys),
        "backup": str(backup),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--list-date", required=True)
    parser.add_argument("--performance-csv", required=True, type=Path)
    parser.add_argument("--v3-ledger", required=True, type=Path)
    parser.add_argument("--v4-ledger", required=True, type=Path)
    parser.add_argument("--backup-dir", required=True, type=Path)
    args = parser.parse_args()
    v3 = read_settlements(args.performance_csv, args.list_date, "V3")
    v4 = read_settlements(args.performance_csv, args.list_date, "V4_SHADOW")
    result = {
        "list_date": args.list_date,
        "performance_csv": str(args.performance_csv),
        "v3": update_ledger(args.v3_ledger, v3, "V3", args.list_date, args.backup_dir / "v3"),
        "v4": update_ledger(args.v4_ledger, v4, "V4", args.list_date, args.backup_dir / "v4"),
        "updated_at": dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).isoformat(),
    }
    print(__import__("json").dumps(result, ensure_ascii=False))
    return 0 if not result["v3"]["unmatched"] and not result["v4"]["unmatched"] else 2


if __name__ == "__main__":
    raise SystemExit(main())
