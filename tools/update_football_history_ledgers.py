"""Append one daily Titan/V3/V4 snapshot to the user-maintained Excel ledgers.

Rows inherit the final existing data-row styles and the append is idempotent:
rerunning one daily snapshot does not append duplicate matches. Prior
settlement history is updated by the dedicated settlement synchronizer.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\codex")


def text(value: object, fallback: str = "") -> str:
    return fallback if value is None or value == "" else str(value)


def number(value: object) -> object:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return value


def kickoff(value: object, list_date: str) -> str:
    raw = text(value).replace("北京时间（东8区）", "").strip()
    if "T" in raw:
        try:
            return dt.datetime.fromisoformat(raw.replace("Z", "+00:00")).astimezone(dt.timezone(dt.timedelta(hours=8))).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            pass
    if " " in raw:
        date_part, time_part = raw.split(" ", 1)
        if len(date_part) <= 5 and "-" in date_part:
            month, day = date_part.split("-", 1)
            return f"{list_date[:4]}-{int(month):02d}-{int(day):02d} {time_part[:5]}:00"
    return raw


def score(row: dict[str, object]) -> str:
    home, away = text(row.get("home_score")), text(row.get("away_score"))
    return f"{home}-{away}" if home and away else "0-0"


def state_label(row: dict[str, object]) -> str:
    state = text(row.get("state"))
    if state == "0":
        return "未开赛"
    if state == "-1":
        return "已结束"
    return "进行中"


def raw_map(path: Path, list_date: str) -> dict[str, dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return {
            str(row.get("match_id", "")): row
            for row in csv.DictReader(handle)
            if str(row.get("list_date", "")) == list_date and row.get("match_id")
        }


def table_last_row(ws) -> int:
    table = next(iter(ws.tables.values()))
    return int("".join(ch for ch in table.ref.split(":")[-1] if ch.isdigit()))


def copy_row_format(ws, source_row: int, destination_row: int) -> None:
    ws.row_dimensions[destination_row].height = ws.row_dimensions[source_row].height
    ws.row_dimensions[destination_row].hidden = ws.row_dimensions[source_row].hidden
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(destination_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        if source.alignment:
            target.alignment = copy(source.alignment)
        if source.protection:
            target.protection = copy(source.protection)


def resize_tables(ws, old_last: int, new_last: int) -> None:
    for table in ws.tables.values():
        start, _ = table.ref.split(":")
        end_ref = table.ref.split(":")[-1]
        end_col = "".join(ch for ch in end_ref if ch.isalpha())
        table.ref = f"{start}:{end_col}{new_last}"


def append_rows(path: Path, rows: list[list[object]], backup_dir: Path) -> int:
    if not rows:
        return 0
    wb = load_workbook(path, data_only=False)
    ws = wb["赛事明细"]
    old_last = table_last_row(ws)
    existing: set[tuple[str, str, str]] = set()
    for row in range(4, old_last + 1):
        existing.add((text(ws.cell(row, 1).value)[:10], text(ws.cell(row, 6).value), text(ws.cell(row, 21).value) if ws.max_column >= 21 else ""))
    pending: list[list[object]] = []
    for values in rows:
        key = (text(values[0])[:10] if values else "", text(values[5]) if len(values) > 5 else "", text(values[20]) if len(values) > 20 else "")
        if key in existing:
            continue
        existing.add(key)
        pending.append(values)
    if not pending:
        wb.close()
        return 0
    backup_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(path, backup_dir / path.name)
    source_row = old_last
    for offset, values in enumerate(pending, start=1):
        target_row = old_last + offset
        copy_row_format(ws, source_row, target_row)
        for col, value in enumerate(values, start=1):
            ws.cell(target_row, col).value = value
    resize_tables(ws, old_last, old_last + len(pending))
    wb.save(path)
    wb.close()
    return len(pending)


def v4_values(raw: dict[str, str], model: dict[str, object], bridge: dict[str, object], source: str) -> list[object]:
    market = model.get("market") or {}
    ah = [raw.get("ah_full_current_home_or_over", ""), raw.get("ah_full_current_line_or_draw", ""), raw.get("ah_full_current_away_or_under", "")]
    v4_status = text(model.get("analysis_status"), "MISSING_DATA")
    grade = text(model.get("grade"))
    reason = ";".join(model.get("reason_codes") or [])
    return [
        text(raw.get("list_date")), text(model.get("kickoff") or raw.get("bj_time")), text(bridge.get("Competition_Domain")),
        text(raw.get("league_tw")), text(raw.get("league_cn")), f"{text(raw.get('home_cn'))} vs {text(raw.get('away_cn'))}",
        number(ah[0]), number(ah[1]), number(ah[2]), "/".join(text(x) for x in ah),
        f"V4-{grade}" if grade else "V4不可用：市场输入缺失", text(model.get("selected_team")), number(model.get("p_giving_cover")),
        "V4 Shadow Only；real_money=false", score(raw), v4_status, "", number(model.get("selected_water_hk")),
        number(model.get("pnl_1u")), source, "", "", "", text(model.get("match_id")), text(model.get("model_id"), "V4"),
        text(model.get("prior_snapshot_id")), grade, number(model.get("p_giving_cover")), number(model.get("p_receiving_cover")),
        number(model.get("fair_handicap")), number(model.get("fair_handicap_low")), number(model.get("fair_handicap_high")), number(model.get("ev_mean")),
        number(model.get("ev_p05")), number(model.get("ev_p10")), number(model.get("ev_p50")), number(model.get("ev_p90")), number(model.get("ev_p95")),
        number(model.get("p_ev_positive")), v4_status, reason, text(model.get("decision_id")), number(model.get("selected_handicap_signed")), number(model.get("selected_water_hk")),
        number(market.get("home_handicap_signed")), number(market.get("away_handicap_signed")), number(market.get("home_water_hk")), number(market.get("away_water_hk")),
        text(model.get("decision_at")), "Titan007/V4", "Daily append；市场输入缺失时不生成决策",
    ]


def v3_values(raw: dict[str, str], model: dict[str, object], source: str) -> list[object]:
    action = text(model.get("action"), "未计算")
    water = number(model.get("water"))
    return [
        text(raw.get("list_date")), kickoff(raw.get("bj_time"), text(raw.get("list_date"))), text(model.get("Competition_Domain")),
        text(raw.get("league_tw")), text(raw.get("league_cn")), f"{text(raw.get('home_cn'))} vs {text(raw.get('away_cn'))}",
        number(raw.get("ah_full_current_home_or_over")), number(raw.get("ah_full_current_line_or_draw")), number(raw.get("ah_full_current_away_or_under")),
        "/".join(text(raw.get(k)) for k in ("ah_full_current_home_or_over", "ah_full_current_line_or_draw", "ah_full_current_away_or_under")),
        text(model.get("direction")), text(model.get("selected_team"), "无，不投"), number(model.get("probability")), action + "；" + text(model.get("reason")),
        score(raw), state_label(raw), "", water, None, source, None, None, None, number(raw.get("ah_full_current_line_or_draw")), water,
        number(raw.get("euro_full_current_home_or_over")), number(raw.get("euro_full_current_line_or_draw")), number(raw.get("euro_full_current_away_or_under")),
        None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None,
        "Titan007/V3", action,
    ]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--list-date", required=True)
    parser.add_argument("--slate-csv", required=True, type=Path)
    parser.add_argument("--v3-bridge", required=True, type=Path)
    parser.add_argument("--v4-json", required=True, type=Path)
    parser.add_argument("--v3-ledger", required=True, type=Path)
    parser.add_argument("--v4-ledger", required=True, type=Path)
    parser.add_argument("--backup-dir", required=True, type=Path)
    args = parser.parse_args()

    raw = raw_map(args.slate_csv, args.list_date)
    bridge_payload = json.loads(args.v3_bridge.read_text(encoding="utf-8"))
    bridge = {str(row.get("match_id")): row for row in bridge_payload.get("matches", [])}
    v4_payload = json.loads(args.v4_json.read_text(encoding="utf-8"))
    v4 = {str(row.get("match_id")): row for row in v4_payload.get("matches", [])}
    source = str(args.slate_csv)

    # The Excel ledgers are bettable-history ledgers, not full roster dumps.
    # Keep only the frozen production V3 actions and V4 A/B/C candidates.
    v4_rows = [v4_values(raw[mid], v4[mid], bridge.get(mid, {}), source) for mid in sorted(raw.keys()) if mid in v4 and text(v4[mid].get("grade")) in {"A", "B", "C"}]
    v3_rows = [v3_values(raw[mid], bridge.get(mid, {}), source) for mid in sorted(raw.keys()) if text(bridge.get(mid, {}).get("action")) in {"可投", "半仓可投"}]
    added_v3 = append_rows(args.v3_ledger, v3_rows, args.backup_dir / "v3")
    added_v4 = append_rows(args.v4_ledger, v4_rows, args.backup_dir / "v4")
    print(json.dumps({"list_date": args.list_date, "raw_rows": len(raw), "v3_rows_appended": added_v3, "v4_rows_appended": added_v4, "backup": str(args.backup_dir)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
