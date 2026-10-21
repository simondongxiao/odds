"""Incremental, lock-aware football ledger storage and dashboard export.

The workbook is the durable audit ledger.  It contains the complete match
population, including unavailable/no-bet rows.  Dashboard JSON is a derived
cache and never acts as the settlement source of truth.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import math
import os
import re
import shutil
import time
import sys
import uuid
from copy import copy
from pathlib import Path
from typing import Any, Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import update_football_history_ledgers as legacy  # noqa: E402
import sync_football_history_settlement as settlement  # noqa: E402


CN_TZ = dt.timezone(dt.timedelta(hours=8))
SHEET_NAME = "赛事明细"
LOCKED_STATUSES = {"已结算", "已开赛", "进行中", "已结束"}
SETTLEMENT_FIELDS = {"赛果", "结算状态", "红黑", "盈亏(u)", "比分来源", "结算水位", "结算盈亏(u)"}
RESULT_LABELS = {"红", "半红", "走", "半黑", "黑"}
META_HEADERS = [
    "Match_ID", "Home_Team", "Away_Team", "Quality_Gate_Passed", "Final_Decision",
    "Row_Lock_State", "Source_Snapshot", "Updated_At", "History_Key",
    "Frozen_Actionable", "Countable_History", "Ledger_Key", "Score_Status",
    "Score_Source_Type", "Score_Fetched_At", "Score_Match_Confidence",
    "Settlement_Audit_Status", "Exclude_Reason",
]
MACHINE_HEADERS = {
    "Ledger_Key", "Score_Status", "Score_Source_Type", "Score_Fetched_At",
    "Score_Match_Confidence", "Settlement_Audit_Status", "Exclude_Reason",
    "History_Key", "Frozen_Actionable", "Countable_History", "Row_Lock_State",
    "Source_Snapshot", "Updated_At", "Quality_Gate_Passed", "Final_Decision",
    "Match_ID", "Home_Team", "Away_Team",
}
DASHBOARD_FIELDS = [
    "日期筛选", "比赛时间", "赛事分布", "赛事类型", "联赛/杯赛", "具体比赛",
    "Match_ID", "Home_Team", "Away_Team", "即时主水", "即时盘口", "即时客水",
    "即时赔率文本", "V4等级/信号", "投注方", "模型有效胜率", "赛果",
    "结算状态", "红黑", "投注水位", "选中盘口", "选中水位", "盈亏(u)", "比分来源",
    "Quality_Gate_Passed", "Final_Decision", "Row_Lock_State",
    "History_Key", "Frozen_Actionable", "Countable_History",
    "Ledger_Key", "Score_Status", "Score_Source_Type", "Score_Fetched_At",
    "Score_Match_Confidence", "Settlement_Audit_Status", "Exclude_Reason",
    "quality_gate_passed", "final_decision", "grade", "analysis_status", "reason_codes", "real_money",
]


class RowCountShrinkError(RuntimeError):
    pass


class ViewReconciliationError(RuntimeError):
    pass


def _clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return value


def _text(value: Any) -> str:
    value = _clean(value)
    return "" if value is None else str(value).strip()


def _atomic_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=_clean) + "\n"
    with tmp.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(data)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(tmp, path)


def _atomic_workbook_save(workbook: Any, path: Path) -> None:
    """Save a valid xlsx to a same-directory temp file, then replace."""
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.stem}.{uuid.uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(tmp)
        with tmp.open("ab") as handle:
            handle.flush()
            os.fsync(handle.fileno())
        last_error = None
        for attempt in range(5):
            try:
                os.replace(tmp, path)
                last_error = None
                break
            except PermissionError as exc:
                last_error = exc
                time.sleep(0.25 * (attempt + 1))
        if last_error is not None:
            raise last_error
    finally:
        tmp.unlink(missing_ok=True)


def _headers(ws: Any) -> dict[str, int]:
    return {_text(ws.cell(3, column).value): column for column in range(1, ws.max_column + 1) if _text(ws.cell(3, column).value)}


def _table_last_row(ws: Any) -> int:
    tables = list(ws.tables.values())
    if not tables:
        return ws.max_row
    return int(re.search(r"\d+$", tables[0].ref.split(":")[-1]).group())


def _resize_tables(ws: Any, last_row: int, last_column: int | None = None) -> None:
    for table in ws.tables.values():
        start = table.ref.split(":", 1)[0]
        end_ref = table.ref.split(":", 1)[1]
        end_column = get_column_letter(last_column) if last_column else re.sub(r"\d", "", end_ref)
        table.ref = f"{start}:{end_column}{last_row}"


def _copy_row_style(ws: Any, source_row: int, destination_row: int, max_column: int) -> None:
    ws.row_dimensions[destination_row].height = ws.row_dimensions[source_row].height
    for column in range(1, max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(destination_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _ensure_metadata_headers(ws: Any, version: str) -> dict[str, int]:
    headers = _headers(ws)
    required = ["Match_ID"] if version == "V3" else []
    required += [header for header in META_HEADERS[1:] if header not in headers]
    column = ws.max_column
    for header in required:
        if header in headers:
            continue
        column += 1
        ws.cell(3, column).value = header
        headers[header] = column
        if header in MACHINE_HEADERS:
            ws.column_dimensions[get_column_letter(column)].hidden = True
    return headers


def _set_machine_header_visibility(ws: Any, headers: dict[str, int]) -> None:
    """Hide only newly-added/audit columns without changing visible layout."""
    for header in MACHINE_HEADERS:
        column = headers.get(header)
        if column:
            ws.column_dimensions[get_column_letter(column)].hidden = True


def _fallback_key(list_date: Any, home: Any, away: Any, match: Any = "") -> str:
    home_text, away_text = _text(home), _text(away)
    if (not home_text or not away_text) and " vs " in _text(match):
        home_text, away_text = [part.strip() for part in _text(match).split(" vs ", 1)]
    return "|".join((_text(list_date)[:10], home_text.casefold(), away_text.casefold()))


def _row_key(ws: Any, row: int, headers: dict[str, int]) -> str:
    match_id = _text(ws.cell(row, headers.get("Match_ID", 0)).value) if headers.get("Match_ID") else ""
    if _valid_match_id(match_id):
        return f"id:{match_id}"
    return "fallback:" + _fallback_key(
        ws.cell(row, headers.get("日期筛选", 1)).value,
        ws.cell(row, headers.get("Home_Team", 0)).value if headers.get("Home_Team") else "",
        ws.cell(row, headers.get("Away_Team", 0)).value if headers.get("Away_Team") else "",
        ws.cell(row, headers.get("具体比赛", 6)).value,
    )


def _incoming_key(raw: dict[str, Any]) -> str:
    match_id = _text(raw.get("match_id"))
    return f"id:{match_id}" if match_id else "fallback:" + _fallback_key(raw.get("list_date"), raw.get("home_cn") or raw.get("home_en"), raw.get("away_cn") or raw.get("away_en"))


def _history_key(competition: Any, match_id: Any, list_date: Any, home: Any, away: Any, match: Any = "") -> str:
    """Stable row identity used by the historical aggregation contract."""
    competition_text = _text(competition) or "未分类"
    match_id_text = _text(match_id)
    if match_id_text and _valid_match_id(match_id_text):
        return f"{competition_text}|id:{match_id_text}"
    fallback = _fallback_key(list_date, home, away, match)
    return f"{competition_text}|{fallback}"


def _ledger_key(version: str, competition: Any, match_id: Any, list_date: Any, home: Any, away: Any, kickoff: Any = "", match: Any = "") -> str:
    """Return the version-scoped UPSERT key used by the durable workbook."""
    mid = _text(match_id)
    if _valid_match_id(mid):
        return f"{version}|{_text(list_date)[:10]}|{mid}"
    competition_text = _text(competition) or "未分类"
    home_text, away_text = _text(home), _text(away)
    if (not home_text or not away_text) and " vs " in _text(match):
        home_text, away_text = [part.strip() for part in _text(match).split(" vs ", 1)]
    return "|".join((version, _text(list_date)[:10], competition_text.casefold(), home_text.casefold(), away_text.casefold(), _text(kickoff)[:16]))


def _valid_match_id(value: Any) -> bool:
    text = _text(value)
    return bool(text) and text.upper() not in {"A", "B", "C", "N", "NA", "N/A", "NONE", "NAN"}


def _weak_history_key(value: Any) -> bool:
    text = _text(value)
    return bool(re.search(r"\|id:(?:A|B|C|N|NA|N/A|NONE|NAN)$", text, flags=re.IGNORECASE))


def _is_actionable_value(value: Any) -> bool:
    return _text(value).upper() in {"BET_GIVING", "BET_RECEIVING", "可投", "半仓可投"}


def _as_bool(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    text = _text(value).lower()
    if text in {"true", "1", "yes", "y", "是", "通过"}:
        return True
    if text in {"false", "0", "no", "n", "否", "未通过"}:
        return False
    return None


def _has_value(value: Any) -> bool:
    return _text(value) not in {"", "None", "nan", "NaN"}


def _frozen_bet_complete(row: dict[str, Any]) -> bool:
    """Whether the original frozen bet fields are sufficient to settle."""
    team = row.get("投注方") or row.get("selected_team") or row.get("Selected_Team")
    line = row.get("即时盘口") or row.get("选中盘口") or row.get("selected_handicap") or row.get("handicap")
    water = row.get("投注水位") or row.get("选中水位") or row.get("selected_water") or row.get("water")
    action = row.get("Final_Decision") or row.get("投注动作") or row.get("投注说明") or row.get("action")
    return _has_value(team) and _has_value(line) and _number(water) is not None and (
        _is_actionable_value(action) or _text(row.get("结算状态")) == "已结算" or _as_bool(row.get("Frozen_Actionable")) is True
    )


def _score_verified(row: dict[str, Any]) -> bool:
    status = _text(row.get("Score_Status"))
    if status in {"SCORE_VERIFIED", "SCORE_AND_BET_VERIFIED"}:
        return True
    score = _text(row.get("赛果")) or _text(row.get("score"))
    return bool(re.fullmatch(r"\d+\s*[-:：]\s*\d+", score)) and _text(row.get("结算状态")) not in {"", "未开赛", "进行中"}


def _countable(row: dict[str, Any]) -> bool:
    explicit = _as_bool(row.get("Countable_History", row.get("countable_history")))
    # Backward compatibility for pre-contract ledgers that predate the hidden
    # audit columns.  Once any contract column is present, the stricter
    # frozen-bet/verified-score checks below apply and an explicit flag cannot
    # bypass them.
    # Frozen_Actionable is derived for every dashboard row, so it is not a
    # reliable marker that the source workbook already has the new contract.
    legacy_contract_absent = not any(
        _has_value(row.get(name))
        for name in ("Score_Status", "Settlement_Audit_Status", "Ledger_Key")
    )
    if legacy_contract_absent:
        return _text(row.get("结算状态")) == "已结算" and _text(row.get("红黑")) in RESULT_LABELS
    # An explicit true flag is accepted only after the frozen fields have been
    # validated.  This prevents a legacy row with a known score from becoming
    # countable merely because its result label was filled first.
    if explicit is True:
        return _text(row.get("结算状态")) == "已结算" and _text(row.get("红黑")) in RESULT_LABELS and _score_verified(row) and _frozen_bet_complete(row)
    if explicit is False:
        return False
    return _text(row.get("结算状态")) == "已结算" and _text(row.get("红黑")) in RESULT_LABELS and _score_verified(row) and _frozen_bet_complete(row)


def _parse_kickoff(value: Any, list_date: str) -> dt.datetime | None:
    normalized = legacy.kickoff(value, list_date)
    try:
        parsed = dt.datetime.fromisoformat(normalized.replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed.replace(tzinfo=CN_TZ) if parsed.tzinfo is None else parsed.astimezone(CN_TZ)


def _locked(ws: Any, row: int, headers: dict[str, int], now: dt.datetime) -> bool:
    status = _text(ws.cell(row, headers.get("结算状态", 16)).value)
    if status in LOCKED_STATUSES or status.startswith("赛果未匹配"):
        return True
    kickoff_value = ws.cell(row, headers.get("比赛时间", 2)).value
    parsed = _parse_kickoff(kickoff_value, _text(ws.cell(row, headers.get("日期筛选", 1)).value)[:10])
    return bool(parsed and parsed <= now)


def _read_csv(path: Path, list_date: str) -> pd.DataFrame:
    frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    frame = frame.loc[frame["list_date"].astype(str).str[:10].eq(list_date)].copy()
    if "match_id" not in frame:
        frame["match_id"] = ""
    frame["_unique_key"] = frame.apply(lambda row: _incoming_key(row.to_dict()), axis=1)
    # The vectorized append contract: latest duplicate wins, but no unique raw
    # match is discarded by a strategy filter.
    return pd.concat([frame.iloc[0:0], frame], ignore_index=True).drop_duplicates(subset=["_unique_key"], keep="last")


def _model_map(path: Path) -> dict[str, dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows = payload.get("matches", payload.get("rows", [])) if isinstance(payload, dict) else payload
    return {_text(row.get("match_id")): row for row in rows if _text(row.get("match_id"))}


def _row_values(version: str, raw: dict[str, Any], model: dict[str, Any], v3_model: dict[str, Any], source: str, headers: dict[str, int]) -> dict[str, Any]:
    if version == "V4":
        values = legacy.v4_values(raw, model, v3_model, source)
    else:
        values = legacy.v3_values(raw, model, source)
    result = {header: values[column - 1] for header, column in headers.items() if column <= len(values)}
    # A missing score must stay empty.  Treating it as 0-0 corrupts settlement.
    if not (_text(raw.get("home_score")) and _text(raw.get("away_score"))):
        result["赛果"] = None
    competition = raw.get("league_cn") or raw.get("league_tw") or raw.get("league_en") or model.get("competition")
    domain = model.get("domain") or model.get("Competition_Domain") or v3_model.get("Competition_Domain") or v3_model.get("domain")
    home = raw.get("home_cn") or raw.get("home_en")
    away = raw.get("away_cn") or raw.get("away_en")
    actionable = model.get("final_decision") if version == "V4" else model.get("action")
    score_ready = bool(_text(raw.get("home_score")) and _text(raw.get("away_score")) and _text(raw.get("state")) == "-1")
    frozen_actionable = _is_actionable_value(actionable)
    score_status = "SCORE_VERIFIED" if score_ready else "SCORE_MISSING"
    settlement_status = _text(result.get("结算状态")) if "结算状态" in result else ""
    result.update({
        "Match_ID": _text(raw.get("match_id")),
        "Home_Team": _text(home),
        "Away_Team": _text(away),
        "Quality_Gate_Passed": model.get("quality_gate_passed") if version == "V4" else None,
        "Final_Decision": actionable,
        "Row_Lock_State": "PREMATCH_MUTABLE",
        "Source_Snapshot": source,
        "Updated_At": dt.datetime.now(CN_TZ).isoformat(),
        "History_Key": _history_key(competition, raw.get("match_id"), raw.get("list_date"), home, away),
        "Ledger_Key": _ledger_key(version, competition, raw.get("match_id"), raw.get("list_date"), home, away, model.get("kickoff") or raw.get("bj_time"), raw.get("match")),
        "Frozen_Actionable": frozen_actionable,
        "Countable_History": False,
        "Score_Status": score_status,
        "Score_Source_Type": "titan007_snapshot" if score_ready else "",
        "Score_Fetched_At": dt.datetime.now(CN_TZ).isoformat() if score_ready else "",
        "Score_Match_Confidence": "exact_match_id" if score_ready and _text(raw.get("match_id")) else ("exact_identity" if score_ready else ""),
        "Settlement_Audit_Status": "UNSETTLED" if not score_ready else "BET_UNVERIFIED",
        "Exclude_Reason": "SCORE_MISSING" if not score_ready else ("NO_BET" if not frozen_actionable else "BET_NOT_SETTLED"),
    })
    if "赛事类型" in result:
        result["赛事类型"] = legacy.type_label(domain, competition)
    return result


def _normalize_competition_type(ws: Any, headers: dict[str, int], start_date: str = "2026-10-08") -> bool:
    """Repair D-column drift without changing frozen market fields.

    Rows from the 10/8 regression either copied E into D or left D blank.  The
    type is derived from the canonical competition name, so every subsequent
    refresh also repairs a stale row before export.
    """
    changed = False
    for row in range(4, _table_last_row(ws) + 1):
        day = _text(ws.cell(row, headers.get("日期筛选", 1)).value)[:10]
        if day < start_date:
            continue
        competition = ws.cell(row, headers.get("联赛/杯赛", 5)).value
        expected = legacy.type_label("", competition)
        cell = ws.cell(row, headers.get("赛事类型", 4))
        if _text(cell.value) != expected:
            cell.value = expected
            changed = True
    return changed


def _normalize_v4_result_styles(ws: Any, headers: dict[str, int]) -> bool:
    """Keep V4 red/black/half/push results visibly distinct."""
    result_col = headers.get("红黑")
    if not result_col:
        return False
    labels = {"红": ("FFF4CCCC", "FFC00000"), "黑": ("FFA6A6A6", "FF000000"), "半红": ("FFFCE4D6", "FFC65911"), "半黑": ("FFD9E1F2", "FF666666"), "走": ("FFFFF2CC", "FF7F6000")}
    changed = False
    for row in range(4, _table_last_row(ws) + 1):
        label = _text(ws.cell(row, result_col).value)
        if label not in labels:
            continue
        fill_color, font_color = labels[label]
        cell = ws.cell(row, result_col)
        if cell.fill.fill_type != "solid" or cell.fill.fgColor.rgb != fill_color:
            cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
            changed = True
        if not cell.font.color or cell.font.color.type != "rgb" or cell.font.color.rgb != font_color:
            old = cell.font
            cell.font = Font(name=old.name, sz=old.sz, b=old.b, i=old.i, u=old.u, strike=old.strike, color=font_color, vertAlign=old.vertAlign, charset=old.charset, family=old.family, scheme=old.scheme)
            changed = True
    return changed


def _cache_records(ws: Any, headers: dict[str, int]) -> list[dict[str, Any]]:
    last = _table_last_row(ws)
    ordered = sorted(headers.items(), key=lambda item: item[1])
    match_id_counts: dict[str, int] = {}
    if headers.get("Match_ID"):
        for row in range(4, last + 1):
            mid = _text(ws.cell(row, headers["Match_ID"]).value)
            if _valid_match_id(mid):
                match_id_counts[mid] = match_id_counts.get(mid, 0) + 1
    records: list[dict[str, Any]] = []
    for row in range(4, last + 1):
        if not any(ws.cell(row, column).value not in (None, "") for _, column in ordered):
            continue
        record = {header: _clean(ws.cell(row, column).value) for header, column in ordered}
        # Backfill metadata for legacy rows without mutating the workbook. This
        # keeps old ledgers countable while the next atomic update persists the
        # new contract columns.
        if not _text(record.get("History_Key")) or _weak_history_key(record.get("History_Key")):
            record["History_Key"] = _history_key(
                record.get("联赛/杯赛"), record.get("Match_ID"), record.get("日期筛选"),
                record.get("Home_Team"), record.get("Away_Team"), record.get("具体比赛"),
            )
        mid = _text(record.get("Match_ID"))
        if not _text(record.get("Ledger_Key")) or match_id_counts.get(mid, 0) > 1:
            record["Ledger_Key"] = _ledger_key(
                "V4" if record.get("Decision_ID") or record.get("V4等级") or record.get("V4等级/信号") else "V3",
                record.get("联赛/杯赛"), "" if match_id_counts.get(mid, 0) > 1 else record.get("Match_ID"), record.get("日期筛选"),
                record.get("Home_Team"), record.get("Away_Team"), record.get("比赛时间"), record.get("具体比赛"),
            )
        frozen = _as_bool(record.get("Frozen_Actionable"))
        if frozen is None:
            frozen = _text(record.get("Row_Lock_State")) == "STARTED_OR_SETTLED_LOCKED" and _is_actionable_value(record.get("Final_Decision") or record.get("投注动作") or record.get("投注说明"))
        record["Frozen_Actionable"] = bool(frozen)
        record["Countable_History"] = _countable(record)
        if not _text(record.get("Exclude_Reason")) and not record["Countable_History"]:
            record["Exclude_Reason"] = "SCORE_MISSING" if not _score_verified(record) else "BET_DIRECTION_MISSING" if not _frozen_bet_complete(record) else "SETTLEMENT_UNRESOLVED"
        records.append(record)
    return records


def _backfill_metadata(ws: Any, headers: dict[str, int], last: int) -> bool:
    """Populate new contract columns without touching frozen market fields."""
    changed = False
    match_id_counts: dict[str, int] = {}
    if headers.get("Match_ID"):
        for scan_row in range(4, last + 1):
            mid = _text(ws.cell(scan_row, headers["Match_ID"]).value)
            if _valid_match_id(mid):
                match_id_counts[mid] = match_id_counts.get(mid, 0) + 1
    for row in range(4, last + 1):
        if not any(ws.cell(row, column).value not in (None, "") for column in range(1, min(ws.max_column, 30) + 1)):
            continue
        values = {header: ws.cell(row, column).value for header, column in headers.items()}
        if headers.get("History_Key") and (not _text(values.get("History_Key")) or _weak_history_key(values.get("History_Key"))):
            values["History_Key"] = _history_key(
                values.get("联赛/杯赛"), "" if match_id_counts.get(_text(values.get("Match_ID")), 0) > 1 else values.get("Match_ID"), values.get("日期筛选"),
                values.get("Home_Team"), values.get("Away_Team"), values.get("具体比赛"),
            )
            ws.cell(row, headers["History_Key"]).value = values["History_Key"]
            changed = True
        if headers.get("Ledger_Key"):
            expected_key = _ledger_key(
                "V4" if headers.get("Decision_ID") or headers.get("V4等级") else "V3",
                values.get("联赛/杯赛"), "" if match_id_counts.get(_text(values.get("Match_ID")), 0) > 1 else values.get("Match_ID"), values.get("日期筛选"),
                values.get("Home_Team"), values.get("Away_Team"), values.get("比赛时间"), values.get("具体比赛"),
            )
            if _text(values.get("Ledger_Key")) != expected_key:
                ws.cell(row, headers["Ledger_Key"]).value = expected_key
                values["Ledger_Key"] = expected_key
                changed = True
        locked = _text(values.get("Row_Lock_State")) == "STARTED_OR_SETTLED_LOCKED" or _text(values.get("结算状态")) in LOCKED_STATUSES
        frozen = _as_bool(values.get("Frozen_Actionable"))
        if frozen is None:
            frozen = locked and _is_actionable_value(values.get("Final_Decision") or values.get("投注动作") or values.get("投注说明"))
            if headers.get("Frozen_Actionable"):
                ws.cell(row, headers["Frozen_Actionable"]).value = bool(frozen)
                values["Frozen_Actionable"] = bool(frozen)
                changed = True
        score_verified = _score_verified(values)
        score_status = "SCORE_VERIFIED" if score_verified else "SCORE_MISSING"
        if headers.get("Score_Status") and _text(values.get("Score_Status")) != score_status:
            ws.cell(row, headers["Score_Status"]).value = score_status
            values["Score_Status"] = score_status
            changed = True
        bet_complete = _frozen_bet_complete(values)
        countable = bool(_text(values.get("结算状态")) == "已结算" and _text(values.get("红黑")) in RESULT_LABELS and score_verified and bet_complete)
        if headers.get("Countable_History") and _as_bool(values.get("Countable_History")) != countable:
            ws.cell(row, headers["Countable_History"]).value = countable
            values["Countable_History"] = countable
            changed = True
        audit = "SETTLEMENT_VERIFIED" if countable else "BET_UNVERIFIED" if score_verified and not bet_complete else "SCORE_PENDING"
        if headers.get("Settlement_Audit_Status") and _text(values.get("Settlement_Audit_Status")) != audit:
            ws.cell(row, headers["Settlement_Audit_Status"]).value = audit
            changed = True
        if headers.get("Exclude_Reason"):
            reason = "" if countable else ("SCORE_MISSING" if not score_verified else "BET_DIRECTION_MISSING" if not bet_complete else "SETTLEMENT_UNRESOLVED")
            if _text(values.get("Exclude_Reason")) != reason:
                ws.cell(row, headers["Exclude_Reason"]).value = reason
                changed = True
    return changed


def update_workbook(
    path: Path,
    version: str,
    raw_frame: pd.DataFrame | None,
    model_map: dict[str, dict[str, Any]] | None,
    v3_map: dict[str, dict[str, Any]] | None,
    source: str,
    backup_dir: Path,
    performance_csv: Path | None = None,
    settlement_date: str = "",
    cache_path: Path | None = None,
) -> dict[str, Any]:
    started = dt.datetime.now(CN_TZ)
    workbook = load_workbook(path, data_only=False)
    ws = workbook[SHEET_NAME]
    headers = _ensure_metadata_headers(ws, version)
    _set_machine_header_visibility(ws, headers)
    last = _table_last_row(ws)
    index = {_row_key(ws, row, headers): row for row in range(4, last + 1)}
    appended = updated = locked = settled = 0
    dirty = False

    if raw_frame is not None:
        for raw in raw_frame.to_dict("records"):
            match_id = _text(raw.get("match_id"))
            model = (model_map or {}).get(match_id, {})
            incoming = _row_values(version, raw, model, (v3_map or {}).get(match_id, {}), source, headers)
            key = _incoming_key(raw)
            fallback_key = "fallback:" + _fallback_key(raw.get("list_date"), raw.get("home_cn") or raw.get("home_en"), raw.get("away_cn") or raw.get("away_en"))
            row = index.get(key) or index.get(fallback_key)
            if row is None:
                last += 1
                _copy_row_style(ws, max(4, last - 1), last, max(ws.max_column, max(headers.values())))
                for header, value in incoming.items():
                    if header in headers:
                        ws.cell(last, headers[header]).value = _clean(value)
                index[key] = last
                index[fallback_key] = last
                appended += 1
                dirty = True
                continue
            is_locked = _locked(ws, row, headers, started)
            if is_locked:
                locked += 1
                if headers.get("Frozen_Actionable") and not _as_bool(ws.cell(row, headers["Frozen_Actionable"]).value):
                    decision = ws.cell(row, headers.get("Final_Decision", 0)).value if headers.get("Final_Decision") else None
                    action = ws.cell(row, headers.get("投注动作", 0)).value if headers.get("投注动作") else None
                    note = ws.cell(row, headers.get("投注说明", 0)).value if headers.get("投注说明") else None
                    if _is_actionable_value(decision or action or note):
                        ws.cell(row, headers["Frozen_Actionable"]).value = True
                        dirty = True
            for header, value in incoming.items():
                if header not in headers:
                    continue
                if is_locked and header not in {"赛果", "Row_Lock_State", "Updated_At"}:
                    # Filling a previously blank technical identity field does
                    # not alter the frozen market decision.
                    if header not in {"Match_ID", "Home_Team", "Away_Team"} or ws.cell(row, headers[header]).value not in (None, ""):
                        continue
                if is_locked and header == "赛果" and value in (None, ""):
                    continue
                target_value = "STARTED_OR_SETTLED_LOCKED" if header == "Row_Lock_State" and is_locked else _clean(value)
                cell = ws.cell(row, headers[header])
                if cell.value != target_value:
                    cell.value = target_value
                    updated += 1
                    dirty = True

    if performance_csv and performance_csv.exists():
        rows = settlement.read_settlements(performance_csv, settlement_date, "V3" if version == "V3" else "V4_SHADOW")
        by_match = {_text(ws.cell(row, headers.get("具体比赛", 6)).value): row for row in range(4, last + 1)}
        by_id = {_text(ws.cell(row, headers["Match_ID"]).value): row for row in range(4, last + 1) if headers.get("Match_ID")}
        for item in rows:
            row = by_id.get(_text(item.get("match_id"))) or by_match.get(_text(item.get("match")))
            if not row:
                continue
            prior_decision = ws.cell(row, headers.get("Final_Decision", 0)).value if headers.get("Final_Decision") else None
            prior_action = ws.cell(row, headers.get("投注动作", 0)).value if headers.get("投注动作") else None
            prior_note = ws.cell(row, headers.get("投注说明", 0)).value if headers.get("投注说明") else None
            current = {header: ws.cell(row, column).value for header, column in headers.items()}
            # The performance CSV is trusted only as a settlement proposal. A
            # score may be recorded without a bet, but a red/black/PnL label
            # is written only when the frozen team, line and water are all
            # present in the historical row.
            score_text = _text(item.get("score"))
            frozen = _is_actionable_value(item.get("action") or prior_decision or prior_action or prior_note) or _as_bool(current.get("Frozen_Actionable")) is True
            bet_complete = _frozen_bet_complete({**current, "投注方": item.get("selected_team") or current.get("投注方"), "即时盘口": item.get("line") or current.get("即时盘口"), "投注水位": item.get("water") or current.get("投注水位"), "Frozen_Actionable": frozen})
            verified_result = item.get("result") in settlement.RESULT_LABEL and bool(score_text) and bet_complete
            changes = {
                "赛果": score_text,
                "结算状态": "已结算" if verified_result else ("赛果已核/冻结投注未核" if score_text else "结算待核"),
                "红黑": settlement.RESULT_LABEL[item["result"]] if verified_result else "",
                "盈亏(u)": legacy.number(item.get("pnl_1u")) if verified_result else None,
                "比分来源": item.get("result_source") or item.get("source"),
                "Row_Lock_State": "STARTED_OR_SETTLED_LOCKED",
                "Updated_At": dt.datetime.now(CN_TZ).isoformat(),
                "Frozen_Actionable": frozen,
                "Countable_History": verified_result,
                "Score_Status": "SCORE_AND_BET_VERIFIED" if verified_result else "SCORE_VERIFIED_BET_UNVERIFIED" if score_text else "SCORE_MISSING",
                "Score_Source_Type": "titan007_snapshot" if "titan" in _text(item.get("result_source") or item.get("source")).lower() else "verified_external" if score_text else "",
                "Score_Fetched_At": dt.datetime.now(CN_TZ).isoformat() if score_text else "",
                "Score_Match_Confidence": "exact_match_id" if _text(item.get("match_id")) else "exact_identity",
                "Settlement_Audit_Status": "SETTLEMENT_VERIFIED" if verified_result else "BET_UNVERIFIED" if score_text else "SCORE_PENDING",
                "Exclude_Reason": "" if verified_result else "BET_DIRECTION_MISSING" if score_text else "SCORE_MISSING",
            }
            changed = False
            for header, value in changes.items():
                if header in headers and ws.cell(row, headers[header]).value != _clean(value):
                    ws.cell(row, headers[header]).value = _clean(value)
                    changed = dirty = True
            settled += int(changed)

    dirty = _normalize_competition_type(ws, headers) or dirty
    if version == "V4":
        dirty = _normalize_v4_result_styles(ws, headers) or dirty
    dirty = _backfill_metadata(ws, headers, last) or dirty
    _resize_tables(ws, last, max(headers.values()))
    records = _cache_records(ws, headers)
    if dirty:
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup = backup_dir / path.name
        if not backup.exists():
            shutil.copy2(path, backup)
        _atomic_workbook_save(workbook, path)
    workbook.close()
    if cache_path:
        _atomic_json(cache_path, {"schema_version": "football-ledger-cache-v1", "version": version, "source_workbook": str(path), "generated_at": dt.datetime.now(CN_TZ).isoformat(), "rows": records})
    return {"version": version, "path": str(path), "input_rows": 0 if raw_frame is None else len(raw_frame), "appended": appended, "cells_updated": updated, "locked_rows": locked, "settled_rows_updated": settled, "ledger_rows": len(records), "atomic_write": dirty, "elapsed_seconds": round((dt.datetime.now(CN_TZ) - started).total_seconds(), 3)}


def _roster_payload(frame: pd.DataFrame, list_date: str) -> dict[str, Any]:
    rows = []
    for item in frame.to_dict("records"):
        rows.append({"key": item["_unique_key"], "match_id": _text(item.get("match_id")), "competition": _text(item.get("league_cn") or item.get("league_tw") or item.get("league_en")), "home": _text(item.get("home_cn") or item.get("home_en")), "away": _text(item.get("away_cn") or item.get("away_en"))})
    counts = pd.Series([row["competition"] for row in rows], dtype="object").value_counts().sort_index().to_dict()
    return {"schema_version": "football-roster-contract-v1", "list_date": list_date, "updated_at": dt.datetime.now(CN_TZ).isoformat(), "total": len(rows), "competition_counts": counts, "matches": rows}


def audit_roster(frame: pd.DataFrame, list_date: str, contract_path: Path, warning_path: Path, exceptions_path: Path | None = None) -> dict[str, Any]:
    incoming = _roster_payload(frame, list_date)
    previous = json.loads(contract_path.read_text(encoding="utf-8")) if contract_path.exists() else None
    if previous:
        old = {item["key"]: item for item in previous.get("matches", [])}
        new = {item["key"]: item for item in incoming["matches"]}
        missing = [old[key] for key in old.keys() - new.keys()]
        exception_rows: list[dict[str, Any]] = []
        if exceptions_path and exceptions_path.exists():
            exception_payload = json.loads(exceptions_path.read_text(encoding="utf-8"))
            exception_rows = exception_payload.get("exceptions", []) if isinstance(exception_payload, dict) else exception_payload
        allowed = {
            _text(item.get("key") or (f"id:{item.get('match_id')}" if item.get("match_id") else "")): item
            for item in exception_rows
            if _text(item.get("status")).upper() in {"CANCELLED", "POSTPONED", "取消", "延期"}
            and _text(item.get("evidence"))
        }
        unauthorized = [item for item in missing if item["key"] not in allowed]
        if unauthorized:
            warning = {"status": "ROW_COUNT_SHRINK_WARNING", "blocked": True, "list_date": list_date, "previous_total": len(old), "incoming_total": len(new), "missing_count": len(missing), "unauthorized_missing_count": len(unauthorized), "missing_matches": unauthorized, "official_exceptions": [allowed[item["key"]] for item in missing if item["key"] in allowed], "message": f"完整赛事矩阵从 {len(old)} 场缩减为 {len(new)} 场；{len(unauthorized)} 场未提供明确取消/延期证明，已阻止覆盖底稿。", "generated_at": dt.datetime.now(CN_TZ).isoformat()}
            _atomic_json(warning_path, warning)
            raise RowCountShrinkError(json.dumps(warning, ensure_ascii=False))
        # New discoveries are appended to the immutable roster contract.
        merged = pd.concat([pd.DataFrame(previous.get("matches", [])), pd.DataFrame(incoming["matches"])], ignore_index=True).drop_duplicates(subset=["key"], keep="last")
        incoming["matches"] = merged.to_dict("records")
        incoming["total"] = len(merged)
        incoming["competition_counts"] = merged["competition"].value_counts().sort_index().to_dict()
    _atomic_json(contract_path, incoming)
    result = {"status": "PASS", "blocked": False, "list_date": list_date, "previous_total": previous.get("total") if previous else None, "incoming_total": len(frame), "contract_total": incoming["total"], "official_exception_count": len(allowed) if previous else 0, "generated_at": dt.datetime.now(CN_TZ).isoformat()}
    _atomic_json(warning_path, result)
    return result


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _number(value: Any) -> float | None:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def _effective_pnl(row: dict[str, Any]) -> float | None:
    direct = _number(row.get("盈亏(u)"))
    if direct is not None:
        return direct
    label = _text(row.get("红黑"))
    water = _number(row.get("结算水位"))
    if water is None:
        water = _number(row.get("投注水位"))
    if label == "红" and water is not None:
        return water
    if label == "半红" and water is not None:
        return water * 0.5
    if label == "走":
        return 0.0
    if label == "半黑":
        return -0.5
    if label == "黑":
        return -1.0
    return None


def _dashboard_row(row: dict[str, Any]) -> dict[str, Any]:
    item = dict(row)
    if not _text(item.get("History_Key", item.get("history_key"))) or _weak_history_key(item.get("History_Key", item.get("history_key"))):
        item["History_Key"] = _history_key(
            item.get("联赛/杯赛"), item.get("Match_ID"), item.get("日期筛选"),
            item.get("Home_Team"), item.get("Away_Team"), item.get("具体比赛"),
        )
    else:
        item["History_Key"] = item.get("History_Key") or item.get("history_key")
    frozen = _as_bool(item.get("Frozen_Actionable", item.get("frozen_actionable")))
    if frozen is None:
        frozen = _text(item.get("Row_Lock_State")) == "STARTED_OR_SETTLED_LOCKED" and _is_actionable_value(item.get("Final_Decision") or item.get("投注动作") or item.get("投注说明"))
    item["Frozen_Actionable"] = bool(frozen)
    item["Countable_History"] = _countable(item)
    if not _text(item.get("Exclude_Reason")) and not item["Countable_History"]:
        item["Exclude_Reason"] = "SCORE_MISSING" if not _score_verified(item) else "BET_DIRECTION_MISSING" if not _frozen_bet_complete(item) else "SETTLEMENT_UNRESOLVED"
    pnl = _effective_pnl(item)
    if isinstance(item.get("盈亏(u)"), str) and item["盈亏(u)"].startswith("="):
        item["盈亏公式"] = item["盈亏(u)"]
    item["盈亏(u)"] = pnl
    return item


def _compact_dashboard_row(row: dict[str, Any]) -> dict[str, Any]:
    item = {field: row.get(field) for field in DASHBOARD_FIELDS if field in row or field in {"Quality_Gate_Passed", "Final_Decision"}}
    item.update({
        "history_key": row.get("History_Key") or row.get("history_key"),
        "ledger_key": row.get("Ledger_Key") or row.get("ledger_key"),
        "score_status": row.get("Score_Status") or row.get("score_status"),
        "settlement_audit_status": row.get("Settlement_Audit_Status") or row.get("settlement_audit_status"),
        "exclude_reason": row.get("Exclude_Reason") or row.get("exclude_reason"),
        "frozen_actionable": bool(_as_bool(row.get("Frozen_Actionable", row.get("frozen_actionable"))) or False),
        "countable_history": bool(_countable(row)),
    })
    return item


def _stats(rows: Iterable[dict[str, Any]]) -> dict[str, Any]:
    items = list(rows)
    settled = [row for row in items if _countable(row) and _text(row.get("红黑")) in RESULT_LABELS]
    counts = {label: sum(_text(row.get("红黑")) == label for row in settled) for label in ["红", "半红", "走", "半黑", "黑"]}
    denominator = counts["红"] + 0.5 * counts["半红"] + counts["黑"] + 0.5 * counts["半黑"]
    pnl = sum(value for value in (_effective_pnl(row) for row in settled) if value is not None)
    return {"total_ledger_rows": len(items), "settled_sample_rows": len(settled), "unsettled_rows": len(items) - len(settled), "result_counts": counts, "effective_win_rate": None if not denominator else round((counts["红"] + 0.5 * counts["半红"]) / denominator, 6), "pnl_u": round(pnl, 4), "roi_per_settled_row": None if not settled else round(pnl / len(settled), 6)}


def _reconcile_views(cache_rows: list[dict[str, Any]], full_payload: dict[str, Any], filtered_payload: dict[str, Any], output_path: Path, workbook_path: Path) -> dict[str, Any]:
    source_keys = {_text(row.get("History_Key") or row.get("history_key")) for row in cache_rows}
    full_rows = full_payload.get("rows", [])
    filtered_rows = filtered_payload.get("rows", [])
    full_keys = {_text(row.get("history_key") or row.get("History_Key")) for row in full_rows}
    filtered_keys = {_text(row.get("history_key") or row.get("History_Key")) for row in filtered_rows}
    mismatches: list[str] = []
    if len(source_keys) != len(cache_rows):
        mismatches.append("EXCEL_DUPLICATE_HISTORY_KEY")
    if len(full_rows) != len(cache_rows):
        mismatches.append("EXCEL_FULL_ROW_COUNT_MISMATCH")
    if source_keys != full_keys:
        mismatches.append("EXCEL_FULL_HISTORY_KEY_MISMATCH")
    if not filtered_keys.issubset(full_keys):
        mismatches.append("FILTERED_NOT_SUBSET_OF_FULL")
    if any(row.get("quality_gate_passed") is not True for row in filtered_rows):
        mismatches.append("FILTERED_GATE_MISMATCH")
    if any(_text(row.get("final_decision")) not in {"BET_GIVING", "BET_RECEIVING"} for row in filtered_rows):
        mismatches.append("FILTERED_ACTION_MISMATCH")
    if any(row.get("frozen_actionable") is not True for row in filtered_rows):
        mismatches.append("FILTERED_FROZEN_ACTIONABLE_MISMATCH")
    expected_stats = {}
    for row in full_rows:
        expected_stats.setdefault(_text(row.get("联赛/杯赛")) or "未分类", []).append(row)
    for competition, rows in expected_stats.items():
        actual = (full_payload.get("competition_stats") or {}).get(competition, {})
        if actual.get("total_ledger_rows") != len(rows) or actual.get("settled_sample_rows") != _stats(rows)["settled_sample_rows"]:
            mismatches.append(f"COMPETITION_STATS_MISMATCH:{competition}")
    audit = {
        "schema_version": "football-three-way-reconciliation-v1",
        "status": "PASS" if not mismatches else "FAIL",
        "source": "excel_cache",
        "excel_rows": len(cache_rows),
        "full_json_rows": len(full_rows),
        "filtered_json_rows": len(filtered_rows),
        "excel_history_key_count": len(source_keys),
        "full_history_key_count": len(full_keys),
        "filtered_subset_of_full": filtered_keys.issubset(full_keys),
        "mismatches": mismatches,
        "workbook": str(workbook_path),
        "workbook_sha256": _sha256(workbook_path),
        "generated_at": dt.datetime.now(CN_TZ).isoformat(),
    }
    _atomic_json(output_path, audit)
    return audit


def export_dashboard(cache_path: Path, v4_json: Path, list_date: str, output_dir: Path, warning_path: Path, workbook_path: Path) -> dict[str, Any]:
    cache = json.loads(cache_path.read_text(encoding="utf-8"))
    ledger_rows = [_dashboard_row(row) for row in cache.get("rows", [])]
    v4_payload = json.loads(v4_json.read_text(encoding="utf-8"))
    models = v4_payload.get("matches", v4_payload.get("rows", []))
    model_by_id = {_text(row.get("match_id")): row for row in models}
    current_rows = [row for row in ledger_rows if _text(row.get("日期筛选"))[:10] == list_date]
    enriched = []
    for row in current_rows:
        item = dict(row)
        model = model_by_id.get(_text(row.get("Match_ID")), {})
        item.update({"quality_gate_passed": model.get("quality_gate_passed"), "final_decision": model.get("final_decision"), "grade": model.get("grade"), "analysis_status": model.get("analysis_status"), "reason_codes": model.get("reason_codes") or [], "real_money": False})
        enriched.append(_compact_dashboard_row(item))
    # The strategy view is the executable/frozen V4 boundary.  Passing the
    # data-quality gate alone is insufficient: neutral and vetoed rows can
    # legitimately pass that gate while still ending as NO_BET.  Keeping
    # those rows here makes the dashboard count disagree with the immutable
    # daily bettable export and can mislabel research rows as actionable.
    filtered_rows = [
        row for row in enriched
        if row.get("quality_gate_passed") is True
        and _text(row.get("final_decision")) in {"BET_GIVING", "BET_RECEIVING"}
        and row.get("frozen_actionable") is True
    ]
    competitions: dict[str, list[dict[str, Any]]] = {}
    for row in ledger_rows:
        competitions.setdefault(_text(row.get("联赛/杯赛")) or "未分类", []).append(row)
    competition_stats = {name: _stats(rows) for name, rows in sorted(competitions.items())}
    full_rows = [_compact_dashboard_row(row) for row in ledger_rows]
    warning = json.loads(warning_path.read_text(encoding="utf-8")) if warning_path.exists() else {"status": "UNKNOWN"}
    common = {"schema_version": "football-dashboard-v2", "list_date": list_date, "generated_at": dt.datetime.now(CN_TZ).isoformat(), "source_workbook": str(workbook_path), "source_workbook_sha256": _sha256(workbook_path), "count_shrink_audit": warning}
    full_payload = {**common, "view": "full_ledger", "label": f"全部赛事底稿 ({len(full_rows)}场)", "summary": _stats(full_rows), "competition_stats": competition_stats, "rows": full_rows}
    filtered_payload = {**common, "view": "v4_strategy", "label": f"V4策略过滤看板 ({len(filtered_rows)}场)", "summary": _stats(filtered_rows), "rows": filtered_rows}
    output_dir.mkdir(parents=True, exist_ok=True)
    full_path, filtered_path = output_dir / "data_full.json", output_dir / "data_filtered.json"
    _atomic_json(full_path, full_payload)
    _atomic_json(filtered_path, filtered_payload)
    reconciliation = _reconcile_views(ledger_rows, full_payload, filtered_payload, output_dir / "three_way_reconciliation.json", workbook_path)
    full_payload["reconciliation"] = reconciliation
    filtered_payload["reconciliation"] = reconciliation
    _atomic_json(full_path, full_payload)
    _atomic_json(filtered_path, filtered_payload)
    _atomic_json(output_dir / "view_manifest.json", {"schema_version": "football-dashboard-manifest-v1", "list_date": list_date, "generated_at": common["generated_at"], "full": {"path": "data_full.json", "rows": len(full_rows)}, "filtered": {"path": "data_filtered.json", "rows": len(filtered_rows)}, "workbook_sha256": common["source_workbook_sha256"], "count_shrink_audit": warning, "reconciliation": reconciliation})
    if reconciliation["status"] != "PASS":
        raise ViewReconciliationError(json.dumps(reconciliation, ensure_ascii=False))
    return {"full_json": str(full_path), "filtered_json": str(filtered_path), "manifest_json": str(output_dir / "view_manifest.json"), "reconciliation_json": str(output_dir / "three_way_reconciliation.json"), "full_rows": len(full_rows), "filtered_rows": len(filtered_rows), "competition_count": len(competition_stats), "reconciliation": reconciliation}


def main() -> int:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    audit = sub.add_parser("audit")
    audit.add_argument("--list-date", required=True); audit.add_argument("--slate-csv", required=True, type=Path); audit.add_argument("--contract", required=True, type=Path); audit.add_argument("--warning", required=True, type=Path); audit.add_argument("--exceptions", type=Path)
    update = sub.add_parser("update")
    update.add_argument("--list-date", required=True); update.add_argument("--slate-csv", required=True, type=Path); update.add_argument("--v3-bridge", required=True, type=Path); update.add_argument("--v4-json", required=True, type=Path); update.add_argument("--v3-ledger", required=True, type=Path); update.add_argument("--v4-ledger", required=True, type=Path); update.add_argument("--backup-dir", required=True, type=Path); update.add_argument("--cache-dir", required=True, type=Path)
    settle = sub.add_parser("settle")
    settle.add_argument("--list-date", required=True); settle.add_argument("--performance-csv", required=True, type=Path); settle.add_argument("--v3-ledger", required=True, type=Path); settle.add_argument("--v4-ledger", required=True, type=Path); settle.add_argument("--backup-dir", required=True, type=Path); settle.add_argument("--cache-dir", required=True, type=Path)
    export = sub.add_parser("export")
    export.add_argument("--list-date", required=True); export.add_argument("--cache", required=True, type=Path); export.add_argument("--v4-json", required=True, type=Path); export.add_argument("--output-dir", required=True, type=Path); export.add_argument("--warning", required=True, type=Path); export.add_argument("--workbook", required=True, type=Path)
    args = parser.parse_args()
    try:
        if args.command == "audit":
            result = audit_roster(_read_csv(args.slate_csv, args.list_date), args.list_date, args.contract, args.warning, args.exceptions)
        elif args.command == "update":
            frame = _read_csv(args.slate_csv, args.list_date); v3 = _model_map(args.v3_bridge); v4 = _model_map(args.v4_json)
            args.cache_dir.mkdir(parents=True, exist_ok=True)
            result = {"list_date": args.list_date, "v3": update_workbook(args.v3_ledger, "V3", frame, v3, v3, str(args.slate_csv), args.backup_dir / "v3", cache_path=args.cache_dir / "ledger_v3.json"), "v4": update_workbook(args.v4_ledger, "V4", frame, v4, v3, str(args.slate_csv), args.backup_dir / "v4", cache_path=args.cache_dir / "ledger_v4.json")}
        elif args.command == "settle":
            args.cache_dir.mkdir(parents=True, exist_ok=True)
            result = {"list_date": args.list_date, "v3": update_workbook(args.v3_ledger, "V3", None, None, None, "", args.backup_dir / "v3", args.performance_csv, args.list_date, args.cache_dir / "ledger_v3.json"), "v4": update_workbook(args.v4_ledger, "V4", None, None, None, "", args.backup_dir / "v4", args.performance_csv, args.list_date, args.cache_dir / "ledger_v4.json")}
        else:
            result = export_dashboard(args.cache, args.v4_json, args.list_date, args.output_dir, args.warning, args.workbook)
    except RowCountShrinkError as exc:
        print(str(exc))
        return 42
    except ViewReconciliationError as exc:
        print(str(exc))
        return 43
    print(json.dumps(result, ensure_ascii=False, default=_clean))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
