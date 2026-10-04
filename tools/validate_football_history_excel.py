"""Validate the Excel ledgers used by the V3/V4 dynamic history HTML.

This is intentionally read-only. It checks the source date, worksheet/table
shape, and the basic formatting continuity at the tail of the data table.
"""

from __future__ import annotations

import argparse
import datetime as dt
from pathlib import Path

from openpyxl import load_workbook


def parse_date(value: object) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = str(value or "").strip()
    if len(text) >= 10:
        try:
            return dt.date.fromisoformat(text[:10])
        except ValueError:
            return None
    return None


def style_signature(ws, row: int) -> tuple[object, ...]:
    cells = [ws.cell(row, col) for col in (1, 2, 3, 5, 17, 19)]
    return tuple(
        (
            cell.style_id,
            cell.number_format,
            cell.alignment.horizontal,
            cell.alignment.vertical,
            cell.alignment.wrap_text,
            cell.font.name,
            cell.font.sz,
            cell.font.bold,
            cell.font.color.type if cell.font.color else None,
            cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None,
            cell.fill.fill_type,
            cell.fill.fgColor.type,
            cell.fill.fgColor.rgb if cell.fill.fgColor.type == "rgb" else None,
            cell.border.left.style,
            cell.border.right.style,
            cell.border.top.style,
            cell.border.bottom.style,
        )
        for cell in cells
    )


def validate(path: Path, target: dt.date, label: str) -> list[str]:
    errors: list[str] = []
    if not path.exists():
        return [f"{label}: Excel 不存在: {path}"]
    try:
        wb_values = load_workbook(path, read_only=False, data_only=True)
        wb_styles = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:  # pragma: no cover - surfaced as a publish blocker
        return [f"{label}: Excel 无法读取: {exc}"]

    if "赛事明细" not in wb_values.sheetnames or "赛事明细" not in wb_styles.sheetnames:
        return [f"{label}: 缺少 赛事明细 工作表"]
    values = wb_values["赛事明细"]
    styles = wb_styles["赛事明细"]
    if not values.tables:
        errors.append(f"{label}: 赛事明细 缺少 Excel Table")
    if values.max_row < 4 or values.max_column < 19:
        errors.append(f"{label}: 赛事明细 表格尺寸异常 {values.max_row}x{values.max_column}")

    dated_rows = [
        row
        for row in range(4, values.max_row + 1)
        if parse_date(values.cell(row, 1).value) is not None
    ]
    dates = [parse_date(values.cell(row, 1).value) for row in dated_rows]
    dates = [value for value in dates if value is not None]
    if not dates:
        errors.append(f"{label}: 没有可识别的列表日期")
    else:
        if max(dates) < target:
            errors.append(f"{label}: 最新日期 {max(dates)} 落后目标日期 {target}")
        if target not in dates:
            errors.append(f"{label}: 未发现目标日期 {target} 的赛事明细")

    if len(dated_rows) >= 2:
        last, previous = dated_rows[-1], dated_rows[-2]
        if styles.row_dimensions[last].height != styles.row_dimensions[previous].height:
            errors.append(f"{label}: 末尾数据行行高与上一数据行不一致")
        if style_signature(styles, last) != style_signature(styles, previous):
            errors.append(f"{label}: 末尾数据行字体/颜色/边框/对齐/数字格式与上一数据行不一致")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--list-date", required=True, help="target list date, YYYY-MM-DD")
    parser.add_argument("--v3-ledger", required=True, type=Path)
    parser.add_argument("--v4-ledger", required=True, type=Path)
    args = parser.parse_args()
    target = dt.date.fromisoformat(args.list_date)
    errors = validate(args.v3_ledger, target, "V3") + validate(args.v4_ledger, target, "V4")
    if errors:
        for error in errors:
            print(error)
        return 1
    print(f"Excel validation OK: target={target} V3={args.v3_ledger.name} V4={args.v4_ledger.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
